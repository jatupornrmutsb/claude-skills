#!/usr/bin/env python3
"""ทดสอบสคริปต์ของ Skill ด้วยไฟล์ตัวอย่างที่สร้างขึ้นเอง (ไม่ใช้ไฟล์จริงของผู้ใช้)

  python selftest.py [--workdir DIR] [--no-render]

ข้อมูลในไฟล์ทดสอบเป็นข้อความสมมุติเพื่อทดสอบเครื่องมือเท่านั้น ห้ามนำไปใช้เป็นผลวิจัย
"""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
PY = sys.executable
RESULTS = []

TH_MIX = ("ระบบต้องรองรับข้อกำหนดระดับ Must และ Should โดยเฉพาะ NFR07 ด้านความปลอดภัยและ DR11 "
          "ด้านข้อมูล ซึ่งเชื่อมโยงกับกรอบ IPOF (Input–Process–Output–Feedback) ตามแนวทาง https://example.org/spec-v2 ")


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail and not cond else ""))


def sh(*args):
    r = subprocess.run([PY, *map(str, args)], capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONUTF8": "1"})
    return r.returncode, r.stdout + r.stderr


def bold_para(doc, text, size=14):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def build_fixture(path, long_body=1, extra_first_cite=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # Letter: ต้องถูกแก้เป็น A4
    sec.left_margin = sec.right_margin = Cm(2.54)
    bold_para(doc, "การออกแบบสถาปัตยกรรมระบบสารสนเทศ: กรณีศึกษาสมมุติสำหรับทดสอบ", 20)
    doc.add_paragraph("ผู้แต่ง ทดสอบ1 และ ผู้แต่ง ทดสอบ2")
    doc.add_paragraph("หน่วยงานสมมุติ อีเมล test@example.org")
    bold_para(doc, "บทคัดย่อ")
    p = doc.add_paragraph()
    r1 = p.add_run("บทความนี้นำเสนอการออกแบบ" + "​" + "สถาปัตยกรรม ")
    r2 = p.add_run(TH_MIX)
    r2._r.get_or_add_rPr().append(_el("w:spacing", val="-10"))
    r2._r.get_or_add_rPr().append(_el("w:w", val="90"))
    r1.font.name = "Angsana New"
    bold_para(doc, "1. บทนำ")
    first = ("การศึกษานี้อ้างอิงแนวคิดเดิม " + ("[10] " if extra_first_cite else "") + "[2] และงานพื้นฐาน [1] "
             "รวมถึงงานกลุ่ม [3]–[5] และ [1, 6] ส่วนแหล่งที่ยังไม่มีในรายการคือ [9] ") + TH_MIX
    doc.add_paragraph(first)
    bold_para(doc, "1.1 วัตถุประสงค์")
    for _ in range(long_body):
        doc.add_paragraph(TH_MIX * 3)
    doc.add_paragraph("This paragraph is entirely in English to check that English-only body text "
                      "gets left-right justification rather than Thai distributed alignment.")
    bold_para(doc, "3. วิธีวิจัย")  # ข้ามหมายเลข 2 โดยเจตนา
    doc.add_paragraph("ตารางแสดงผลสมมุติ")
    bold_para(doc, "ตารางที่ 1 ข้อกำหนดตัวอย่าง (ข้อมูลสมมุติ)")
    t = doc.add_table(rows=22, cols=4)
    t.style = "Table Grid"
    hdr = ["รหัส", "ข้อกำหนด", "MoSCoW", "ค่าเฉลี่ย"]
    for c, h in enumerate(hdr):
        t.cell(0, c).text = h
    for rr in range(1, 22):
        t.cell(rr, 0).text = f"NFR{rr:02d}"
        t.cell(rr, 1).text = "ระบบต้องบันทึกประวัติการเข้าถึงข้อมูลและรองรับการตรวจสอบย้อนหลัง"
        t.cell(rr, 2).text = "Must"
        t.cell(rr, 3).text = f"4.{rr % 10}"
        trPr = t.rows[rr]._tr.get_or_add_trPr()
        trPr.append(_el("w:trHeight", val="300", hRule="exact"))
    a = t.cell(5, 2).merge(t.cell(6, 2))  # ผสานแนวตั้ง
    a.text = "Should"
    t.cell(8, 1).merge(t.cell(8, 2))  # ผสานแนวนอน
    for row in t.rows:  # Table Grid ไม่กำหนด grid เอง: ตั้ง 4 คอลัมน์เท่ากัน
        for cell in row.cells:
            cell.width = Cm(3.0)
    bold_para(doc, "เอกสารอ้างอิง")
    refs = [
        "A. Author, “Paper one,” IEEE Access, vol. 1, pp. 1–10, 2020, doi: 10.1109/ACCESS.2020.0000001.",
        "B. Author, “Paper two,” in Proc. Conf., 2021, pp. 11–20.",
        "C. Author, Book Three. New York, NY, USA: Pub, 2019.",
        "D. Author, “Report four,” Org, Tech. Rep., 2022.",
        "E. Author, “Web five.” Example. https://example.org/five (accessed Jan. 1, 2024).",
        "F. Author, “Paper six,” J. Test, vol. 6, no. 2, pp. 1–2, 2023.",
        "G. Author, “Unused seven,” J. Test, 2018.",
        "H. Author, “Unused eight,” J. Test, 2017.",
    ]
    if extra_first_cite:
        refs += ["", "I. Author, “New ten,” J. Test, 2024."]
        refs[8] = "X. Author, “Nine,” J. Test, 2016."
    for i, rtxt in enumerate(refs, 1):
        doc.add_paragraph(f"[{i}] {rtxt}")
    bold_para(doc, "ภาคผนวก ก")
    doc.add_paragraph("รายละเอียดเพิ่มเติมของภาคผนวกอ้างถึง [4] และข้อมูลประกอบ")
    doc.core_properties.author = "Real Author Name"
    doc.save(path)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def audit_json(path, *extra):
    out = Path(str(path) + ".audit.json")
    sh(HERE / "docx_audit.py", path, "--json", out, *extra)
    return json.loads(out.read_text(encoding="utf8"))


def codes(a):
    return {i["code"] for i in a["issues"]}


def hidden_count(path):
    d = Document(path)
    return sum(t.text.count("​") for t in d.element.body.iter(qn("w:t")) if t.text)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    wd = Path(a.workdir or tempfile.mkdtemp(prefix="jdtaia_selftest_"))
    wd.mkdir(parents=True, exist_ok=True)
    src = wd / "fixture.docx"
    build_fixture(src)
    before = audit_json(src)
    check("audit: พบหน้าไม่ใช่ A4", "PAGE_SIZE" in codes(before))
    check("audit: พบ U+200B", "HIDDEN_CHAR" in codes(before))
    check("audit: พบระยะ/สัดส่วนตัวอักษรผิดปกติ", {"CHAR_SPACING", "CHAR_SCALE"} <= codes(before))
    check("audit: พบเลขหัวข้อข้าม 1.1→3", "HEADING_SEQ" in codes(before))
    check("audit: พบ “ตารางแสดง”", "CAPTION_WORDING" in codes(before))
    check("audit: พบตารางยาวไม่ทำซ้ำหัว + แถวสูง exact", {"TABLE_NO_REPEAT_HEADER", "TABLE_EXACT_HEIGHT"} <= codes(before))
    check("audit: พบ metadata ผู้แต่ง", "META_AUTHOR" in codes(before))
    check("audit: พบตัวหนา/ขนาดที่ไม่มีผลกับอักษรไทย (bCs/szCs)", {"BCS_MISSING", "SZCS_MISSING"} <= codes(before))
    check("audit: นับสถานะก่อนแก้ (ตาราง 1, อ้างอิง 8)", before["state"]["tables"] == 1 and before["state"]["reference_entries"] == 8,
          json.dumps(before["state"], ensure_ascii=False))

    # 1) จัดรูปแบบอย่างเดียว: ข้อความ/ตัวเลขไม่เปลี่ยน และไม่เพิ่มอักขระแฝง
    fmt = wd / "fmt.docx"
    rc, out = sh(HERE / "apply_layout.py", src, fmt, "--page", "--columns", "2", "--split-before", "1. บทนำ",
                 "--fonts", "--align", "--spacing", "--reset-char", "--tables", "--anonymize-metadata")
    check("apply_layout ทำงาน", rc == 0, out[-400:])
    rc, out = sh(HERE / "text_diff.py", src, fmt)
    check("format-only: text_diff ผ่าน (ไม่มี WORDING/ตัวเลขเปลี่ยน)", rc == 0, out[-600:])
    check("format-only: ไม่เพิ่ม U+200B", hidden_count(fmt) == hidden_count(src))
    after = audit_json(fmt, "--columns", "2")
    check("หลังจัด: A4 และขอบถูกต้อง", not ({"PAGE_SIZE", "MARGIN"} & codes(after)))
    check("หลังจัด: section [1,2] คอลัมน์", [s["columns"] for s in after["sections"]] == [1, 2], str([s["columns"] for s in after["sections"]]))
    check("หลังจัด: ไม่มีระยะ/สัดส่วนตัวอักษรผิดปกติ", not ({"CHAR_SPACING", "CHAR_SCALE"} & codes(after)))
    check("หลังจัด: ฟอนต์ครบทุกช่อง", not ({"FONT", "FONT_CS_MISSING"} & codes(after)))
    check("หลังจัด: เติม bCs/szCs ให้อักษรไทย", not ({"BCS_MISSING", "SZCS_MISSING"} & codes(after)))
    check("หลังจัด: เนื้อความไทยเป็น thaiDistribute", "NOT_THAI_DIST" not in codes(after))
    check("หลังจัด: หัวข้อไม่ถูกกระจาย", "DIST_ON_HEADING" not in codes(after))
    check("หลังจัด: ทำซ้ำหัวตาราง และเลิก exact", not ({"TABLE_NO_REPEAT_HEADER", "TABLE_EXACT_HEIGHT"} & codes(after)))
    check("หลังจัด: ยังรายงานเลขหัวข้อข้าม (ไม่แก้เนื้อหาเอง)", "HEADING_SEQ" in codes(after))
    check("หลังจัด: ยังรายงาน U+200B เดิม (ไม่ลบเองโดยไม่สั่ง)", "HIDDEN_CHAR" in codes(after))
    check("หลังจัด: ล้าง metadata ผู้แต่ง", "META_AUTHOR" not in codes(after))
    eng = [i for i in after["issues"] if i["code"] == "EN_THAI_DIST"]
    check("หลังจัด: อังกฤษล้วนไม่ใช้ thaiDistribute", not eng)
    check("ตารางเซลล์ผสานคงอยู่", after["tables"][0]["vMerge"] >= 2 and after["tables"][0]["gridSpan"] >= 1, str(after["tables"][0]))

    # 1b) ลบอักขระแฝงเมื่อสั่ง → HIDDEN_ONLY เท่านั้น
    clean = wd / "clean.docx"
    sh(HERE / "apply_layout.py", fmt, clean, "--strip-hidden")
    rc, out = sh(HERE / "text_diff.py", fmt, clean)
    check("strip-hidden: ความต่างเป็น HIDDEN_ONLY และผ่าน", rc == 0 and "HIDDEN_ONLY" in out and "WORDING" not in out.split("สรุปความต่าง:")[1].split("\n")[0], out[-300:])
    check("strip-hidden: ไม่เหลือ U+200B", hidden_count(clean) == 0)

    # 1c) text_diff จับการเปลี่ยนถ้อยคำ/ตัวเลขได้จริง
    tam = wd / "tampered.docx"
    d = Document(clean)
    for p in d.paragraphs:
        if p.text.startswith("This paragraph"):
            p.runs[0].text = p.runs[0].text.replace("entirely", "mostly")
    d.tables[0].cell(3, 3).paragraphs[0].runs[0].text = "9.9"
    d.save(tam)
    rc, out = sh(HERE / "text_diff.py", clean, tam)
    check("text_diff จับถ้อยคำและตัวเลขที่เปลี่ยน", rc == 1 and "WORDING" in out and "9.9" in out, out[-400:])
    brk = wd / "break.docx"
    d = Document(clean)
    for p in d.paragraphs:
        if p.text.startswith("การศึกษานี้"):
            r = p.runs[0]
            r._r.append(OxmlElement("w:br"))
    d.save(brk)
    rc, out = sh(HERE / "text_diff.py", clean, brk)
    check("text_diff แยกการขึ้นบรรทัด (BREAK_ONLY) ออกจากถ้อยคำ", rc == 0 and "BREAK_ONLY" in out, out[-300:])

    # 2) สลับ 2 ↔ 1 คอลัมน์ ไม่เหลือคอลัมน์เดิม
    one = wd / "one_col.docx"
    sh(HERE / "apply_layout.py", fmt, one, "--columns", "1")
    a1 = audit_json(one, "--columns", "1")
    check("2→1 คอลัมน์: ทุก section เป็น 1", all(s["columns"] == 1 for s in a1["sections"]) and "COLS_LEFTOVER" not in codes(a1))
    two = wd / "two_col.docx"
    sh(HERE / "apply_layout.py", one, two, "--columns", "2", "--keep-first-single")
    a2 = audit_json(two, "--columns", "2")
    check("1→2 คอลัมน์: section แรก 1 / เนื้อหา 2", [s["columns"] for s in a2["sections"]] == [1, 2])
    rc, _ = sh(HERE / "text_diff.py", src, two)
    check("สลับคอลัมน์ไปมา: ข้อความไม่เปลี่ยน", rc == 0)

    # 3) อ้างอิง IEEE
    rc, out = sh(HERE / "ieee_refs.py", fmt)
    check("refs: จับลำดับอ้างครั้งแรกผิด", "ลำดับการอ้างครั้งแรกไม่ใช่" in out, out)
    check("refs: จับ [9] ไม่มีในรายการ", "ไม่มีในรายการ: [9]" in out, out)
    check("refs: จับรายการไม่ถูกอ้าง [7, 8]", "ไม่ถูกอ้าง: [7, 8]" in out, out)
    check("refs: เตือนวันที่เข้าถึงต้องเป็นวันที่ตรวจจริง", "วันที่เข้าถึง" in out)
    # เพิ่มแหล่งใหม่ [10] ที่อ้างก่อนทุกแหล่ง แล้วเรียงใหม่
    add = wd / "add_ref.docx"
    build_fixture(add, extra_first_cite=True)
    ren = wd / "renum.docx"
    rc, out = sh(HERE / "ieee_refs.py", add, "--renumber", ren)
    check("renumber ทำงาน", rc == 0, out[-400:])
    rc, out2 = sh(HERE / "ieee_refs.py", ren)
    check("หลัง renumber: ลำดับอ้างครั้งแรกเป็น 1..n", "ลำดับการอ้างครั้งแรกไม่ใช่" not in out2 and "ไม่เรียง" not in out2, out2)
    d = Document(ren)
    body = "\n".join(p.text for p in d.paragraphs)
    check("หลัง renumber: แหล่งใหม่ [10] กลายเป็น [1] และรายการ [1] คือ New ten",
          "การศึกษานี้อ้างอิงแนวคิดเดิม [1] [2]" in body and "[1] I. Author" in body, body[:300])
    check("หลัง renumber: ช่วง [3]–[5] ยังคงรูปแบบช่วง", "[4]–[6]" in body, body[:400])
    check("หลัง renumber: ภาคผนวกอ้าง [4] เดิม → [5]", "ภาคผนวกอ้างถึง [5]" in body)
    rc, out = sh(HERE / "text_diff.py", add, ren)
    check("text_diff: เปลี่ยนเลขอ้างอิงโดยไม่อนุญาต = ไม่ผ่าน", rc == 1 and ("RENUMBER" in out or "REORDER" in out), out[-300:])
    rc, out = sh(HERE / "text_diff.py", add, ren, "--allow-renumber")
    check("text_diff: อนุญาต renumber แล้วไม่มี WORDING (รายงาน RENUMBER/REORDER แยก)", rc == 0 and "REORDER" in out, out[-700:])

    # 3b) โครงสร้าง XML ไม่เสียหาย (ใช้ validate.py ของ skill docx ถ้ามี)
    val = Path("/mnt/skills/public/docx/scripts/office/validate.py")
    if val.exists():
        for f, orig in ((fmt, src), (two, src), (ren, add)):
            r = subprocess.run([PY, str(val), str(f), "--original", str(orig)], capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONUTF8": "1"}, cwd=val.parent)
            check(f"schema: {f.name} ไม่มี error ใหม่", "PASSED" in r.stdout, r.stdout[-300:])

    # 4) เข้าถึงไฟล์ไม่ได้
    rc, out = sh(HERE / "docx_audit.py", wd / "missing" / "nofile.docx")
    check("ไฟล์หาย: แจ้งให้ค้นสำเนาก่อนขอซ้ำ", rc != 0 and "ค้นสำเนา" in out, out[-200:])
    rc, out = sh(HERE / "apply_layout.py", src, src, "--page")
    check("ห้ามเขียนทับต้นฉบับ", rc != 0 and "ห้ามเขียนทับ" in out)

    # 5) เรนเดอร์ + จำนวนหน้าเกินเกณฑ์
    if not a.no_render:
        rd = wd / "render"
        rc, out = sh(HERE / "render_check.py", fmt, "--outdir", rd, "--png", "--json", rd / "r.json")
        check("render_check ทำงานและสร้าง PNG", (rd / "r.json").exists() and json.loads((rd / "r.json").read_text())["png"], out[-400:])
        r = json.loads((rd / "r.json").read_text())
        check("render: พบหน้าเริ่มเอกสารอ้างอิงและภาคผนวก", r["pages"]["references_start_page"] and r["pages"]["appendix_start_page"], json.dumps(r["pages"], ensure_ascii=False))
        longf = wd / "long.docx"
        build_fixture(longf, long_body=60)
        longfmt = wd / "long_fmt.docx"
        sh(HERE / "apply_layout.py", longf, longfmt, "--page", "--columns", "2", "--split-before", "1. บทนำ", "--fonts", "--align", "--spacing")
        rc, out = sh(HERE / "render_check.py", longfmt, "--outdir", wd / "render_long", "--json", wd / "render_long" / "r.json")
        rl = json.loads((wd / "render_long" / "r.json").read_text())
        check("หน้าเกินเกณฑ์: รายงาน 'ไม่ผ่าน' และ exit≠0 (ไม่ย่ออัตโนมัติ)", rc != 0 and not rl["within_limit"] and rl["counted_pages"] > 10,
              f"pages={rl['counted_pages']}")
        rc, out = sh(HERE / "text_diff.py", longf, longfmt)
        check("หน้าเกินเกณฑ์: เนื้อหาไม่ถูกตัด", rc == 0)
        check("render: เตือนเมื่อไม่มีฟอนต์ TH Sarabun New", r["fonts"]["th_sarabun_installed"] or "TH Sarabun New ไม่ได้ใช้จริง" in out or not rl["fonts"]["th_sarabun_installed"])

    ok = sum(1 for _, c, _ in RESULTS if c)
    print(f"\nสรุป: ผ่าน {ok}/{len(RESULTS)}  (workdir: {wd})")
    return 0 if ok == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
