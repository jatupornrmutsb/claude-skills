"""ฟังก์ชันกลางสำหรับสคริปต์ของ Skill jdtaia-ieee-document

อ่าน .docx ผ่าน python-docx + lxml โดยไม่แก้ข้อความ
ใช้ได้กับ Python 3.9+ ต้องมี python-docx (pip install python-docx)
"""
from __future__ import annotations

import re
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
import unicodedata
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn

THAI_RE = re.compile(r"[฀-๿]")
LATIN_RE = re.compile(r"[A-Za-z]")

# อักขระแฝงที่ห้ามใช้แก้ช่องไฟ (ตรวจและรายงานเท่านั้น)
HIDDEN_CHARS = {
    "​": "ZERO WIDTH SPACE (U+200B)",
    "‌": "ZERO WIDTH NON-JOINER (U+200C)",
    "‍": "ZERO WIDTH JOINER (U+200D)",
    "⁠": "WORD JOINER (U+2060)",
    "﻿": "ZERO WIDTH NO-BREAK SPACE / BOM (U+FEFF)",
    "­": "SOFT HYPHEN (U+00AD)",
    "‎": "LEFT-TO-RIGHT MARK (U+200E)",
    "‏": "RIGHT-TO-LEFT MARK (U+200F)",
}

# ค่าตั้งต้น JDTAIA (twips; 1 ซม. = 567 twips โดยประมาณ)
CM = 566.929
SPEC = {
    "page_w": 11906, "page_h": 16838,  # A4 portrait
    "top": round(3.5 * CM), "bottom": round(3.0 * CM),
    "left": round(3.0 * CM), "right": round(2.5 * CM),
    "text_width": round(15.5 * CM),
    "col_width": round(7.5 * CM), "col_space": round(0.5 * CM),
    "font": "TH Sarabun New",
}
TOL = 20  # twips ที่ยอมให้คลาดจากการปัดเศษ (~0.035 ซม.)

CAPTION_TABLE_RE = re.compile(r"^\s*ตารางที่\s*(\d+)")
CAPTION_FIG_RE = re.compile(r"^\s*รูปที่\s*(\d+)")
BAD_TABLE_CAPTION_RE = re.compile(r"^\s*ตารางแสดง")
HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,3})\.?\s+\S")
REF_HEADING_RE = re.compile(r"^\s*(\d+\.?\s*)?(เอกสารอ้างอิง|References|REFERENCES|บรรณานุกรม)\s*$")
APPENDIX_HEADING_RE = re.compile(r"^\s*(ภาคผนวก|Appendix|APPENDIX)\b")


def load(path: str):
    p = __import__("pathlib").Path(path)
    if not p.exists():
        raise SystemExit(f"เข้าถึงไฟล์ไม่ได้: {path}\n→ ค้นสำเนาในพื้นที่ทำงาน/ไฟล์แนบ/แหล่งที่เชื่อมต่อก่อนขอไฟล์จากผู้ใช้ซ้ำ"
                         " (ดู references/qa-delivery.md หัวข้อ 'เมื่อเข้าถึงไฟล์ไม่ได้')")
    if p.suffix.lower() != ".docx":
        raise SystemExit(f"รองรับเฉพาะ .docx: {path} (ถ้าเป็น .doc ให้แปลงเป็น .docx ในสำเนาก่อน)")
    return Document(str(p))


def body_items(doc) -> Iterator[tuple[str, object]]:
    """เดินตามลำดับในเนื้อเอกสาร: ('p', paragraph) หรือ ('t', table)"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield "t", Table(child, doc)


def all_paragraphs(doc):
    """ย่อหน้าทั้งหมดในเนื้อหา รวมในตาราง (ซ้อนได้) พร้อมป้ายตำแหน่ง"""
    out = []
    for i, (kind, obj) in enumerate(body_items(doc)):
        if kind == "p":
            out.append((f"P{i}", obj, False))
        else:
            for loc, p in table_paragraphs(obj, f"T{i}"):
                out.append((loc, p, True))
    return out


def table_paragraphs(table, prefix):
    seen = set()
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if cell._tc in seen:  # เซลล์ผสานแนวนอนคืนค่าซ้ำ
                continue
            seen.add(cell._tc)
            for k, p in enumerate(cell.paragraphs):
                yield f"{prefix}.R{r}C{c}.p{k}", p
            for t2, inner in enumerate(cell.tables):
                yield from table_paragraphs(inner, f"{prefix}.R{r}C{c}.T{t2}")


def para_text(p) -> str:
    """ข้อความของย่อหน้า: <w:br/> = \n, <w:tab/> = \t (รวม run ใน hyperlink)"""
    parts = []
    for el in p._p.iter():
        tag = el.tag
        if tag == qn("w:t"):
            parts.append(el.text or "")
        elif tag == qn("w:tab") and el.getparent().tag == qn("w:r"):
            parts.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        elif tag == qn("w:noBreakHyphen"):
            parts.append("‑")
    return "".join(parts)


def has_thai(s: str) -> bool:
    return bool(THAI_RE.search(s))


def has_latin(s: str) -> bool:
    return bool(LATIN_RE.search(s))


def normalize(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def style_chain(p):
    names = []
    st = p.style
    while st is not None:
        names.append((st.name or "").lower())
        st = st.base_style
    return names


def para_jc(p) -> str | None:
    pPr = p._p.pPr
    if pPr is not None:
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            return jc.get(qn("w:val"))
    st = p.style
    while st is not None:
        spPr = st.element.find(qn("w:pPr"))
        if spPr is not None and spPr.find(qn("w:jc")) is not None:
            return spPr.find(qn("w:jc")).get(qn("w:val"))
        st = st.base_style
    return None


def is_bold(p) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    if not runs:
        return False
    def rb(r):
        rPr = r._r.rPr
        if rPr is not None:
            for t in ("w:b", "w:bCs"):
                e = rPr.find(qn(t))
                if e is not None and e.get(qn("w:val"), "true") not in ("0", "false"):
                    return True
        return bool(r.style is not None and r.style.font.bold)
    return all(rb(r) for r in runs) or bool(p.style.font.bold)


def classify(p, text: str, in_table: bool, index_hint: int | None = None) -> str:
    """จัดประเภทย่อหน้าแบบประมาณการ เพื่อเลือกการจัดแนว
    คืนค่า: empty, title, heading, caption_table, caption_figure, bad_caption,
    table_cell, reference, equation, body
    ผลเป็นเพียงข้อเสนอ ต้องให้ผู้ใช้/ผู้ตรวจยืนยันในกรณีก้ำกึ่ง"""
    t = text.strip()
    if not t:
        if p._p.find(".//" + qn("m:oMath")) is not None:
            return "equation"
        return "empty"
    if p._p.find(".//" + qn("m:oMath")) is not None and len(t) < 40:
        return "equation"
    if in_table:
        return "table_cell"
    chain = style_chain(p)
    if any(n == "title" for n in chain):
        return "title"
    if any(n.startswith("heading") or "หัวเรื่อง" in n for n in chain):
        return "heading"
    if CAPTION_TABLE_RE.match(t):
        return "caption_table"
    if CAPTION_FIG_RE.match(t):
        return "caption_figure"
    if BAD_TABLE_CAPTION_RE.match(t):
        return "bad_caption"
    if any(n == "caption" or "คำบรรยาย" in n for n in chain):
        return "caption_figure"
    if REF_HEADING_RE.match(t) or (APPENDIX_HEADING_RE.match(t) and len(t) <= 80):
        return "heading"
    if re.match(r"^\s*\[\d+\]\s", t):
        return "reference"
    if re.match(r"^\s*(บทคัดย่อ|Abstract|ABSTRACT|คำสำคัญ|Keywords?)\s*[:：]?\s*$", t):
        return "heading"
    if HEADING_RE.match(t) and len(t) <= 90 and is_bold(p):
        return "heading"
    return "body"


def sections(doc):
    return list(doc.sections)


def section_cols(sec):
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is None:
        return 1, None
    return int(cols.get(qn("w:num"), "1")), cols.get(qn("w:space"))


def twips_to_cm(v) -> float:
    return round(int(v) / CM, 2)
