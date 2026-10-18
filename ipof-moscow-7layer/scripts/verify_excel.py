#!/usr/bin/env python3
"""recalculate ไฟล์ Excel แล้วตรวจ (1) ไม่มี error (2) ตัวเลขตรงกับการคำนวณแยกด้วย Python

  python verify_excel.py สรุปผล_V1.xlsx [--data DIR] [--keep-recalc]

ใช้ <xlsx>.manifest.json ที่ build_excel.py เขียนไว้ และ scores.csv / choices.csv เดิม
recalc ด้วย LibreOffice (headless) ลงสำเนาชั่วคราว ไม่แก้ไฟล์ส่งมอบ
ออกด้วยรหัส 1 เมื่อพบ error หรือค่าไม่ตรง
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import shutil
import subprocess
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ipof_lib as L  # noqa: E402

ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")


def recalc(xlsx, outdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("ไม่พบ LibreOffice — recalculate ไม่ได้ ห้ามรับรองว่าสูตรไม่มี error")
    prof = tempfile.mkdtemp(prefix="lo_")
    subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(outdir), str(xlsx)], capture_output=True, text=True, timeout=240,
                   env=dict(os.environ, HOME=prof))
    shutil.rmtree(prof, ignore_errors=True)
    out = Path(outdir) / Path(xlsx).name
    if not out.exists():
        raise SystemExit("recalculate ไม่สำเร็จ")
    return out


def close(a, b, tol=0.005):
    if a is None or b is None:
        return a == b
    if isinstance(b, float) and math.isnan(b):
        return a in (None, "") or (isinstance(a, str) and a.startswith("#DIV"))
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx")
    ap.add_argument("--data")
    ap.add_argument("--keep-recalc", action="store_true")
    a = ap.parse_args(argv)
    man = json.loads(Path(a.xlsx + ".manifest.json").read_text(encoding="utf-8"))
    data = Path(a.data or man["data"])
    tmp = Path(tempfile.mkdtemp(prefix="recalc_"))
    rc = recalc(a.xlsx, tmp)
    wb = load_workbook(rc, data_only=True)
    errs, nform = [], 0
    wbf = load_workbook(a.xlsx)
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nform += 1
                    v = wb[ws.title][c.coordinate].value
                    if v is None or (isinstance(v, str) and v.startswith(ERRORS)):
                        if v is None and ('""' in c.value or "STDEV" in c.value):
                            continue  # สูตรที่ตั้งใจคืนค่าว่าง หรือ STDEV ของข้อมูล 1 ค่า
                        errs.append(f"{ws.title}!{c.coordinate} = {v!r} ({c.value[:60]})")
    scores = L.load_scores(data / "scores.csv")
    choices = []
    if (data / "choices.csv").exists():
        with open(data / "choices.csv", encoding="utf-8-sig") as f:
            choices = list(csv.DictReader(f))
    mism = []
    for m in man["rows"]:
        ws = wb[m["sheet"]]
        r = m["row"]
        if m["kind"] in ("item", "pooled"):
            st = L.pooled(scores, m["items"])
            got_m, got_s = ws.cell(r, 4).value, ws.cell(r, 5).value
            if not close(got_m, st["mean"]) or (st["n_values"] > 1 and not close(got_s, st["sd"])):
                mism.append(f"{m['sheet']} แถว {r} {m.get('label', m['items'])}: Excel {got_m}/{got_s} ≠ Python {st['mean']:.4f}/{st['sd']:.4f}")
            want_lvl = L.interpret(st["mean"])
            if ws.cell(r, 6).value != want_lvl:
                mism.append(f"{m['sheet']} แถว {r}: แปลผล {ws.cell(r, 6).value} ≠ {want_lvl}")
        elif m["kind"] == "choice":
            n = sum(1 for c in choices if c["question"] == m["question"] and c["option"] == m["option"] and c["checked"] == "1")
            if ws.cell(r, 3).value != n:
                mism.append(f"{m['sheet']} แถว {r}: จำนวน {ws.cell(r, 3).value} ≠ {n}")
            pct = n / man["k"] * 100 if man["k"] else 0
            if not close(ws.cell(r, 4).value, pct):
                mism.append(f"{m['sheet']} แถว {r}: ร้อยละ {ws.cell(r, 4).value} ≠ {pct:.2f}")
        elif m["kind"] == "agree":
            code = m["items"][0]
            n = sum(1 for v in scores.get(f"{code}#agree", {}).values() if v == "agree")
            if ws.cell(r, m["agree_col"]).value != n:
                mism.append(f"{m['sheet']} {code}: เห็นชอบ {ws.cell(r, m['agree_col']).value} ≠ {n}")
    print(f"สูตร {nform} ช่อง, error {len(errs)}, ตรวจตัวเลขแยก {len(man['rows'])} แถว, ไม่ตรง {len(mism)}")
    for e in errs[:30]:
        print("  [error]", e)
    for e in mism[:30]:
        print("  [ไม่ตรง]", e)
    if a.keep_recalc:
        print("สำเนาที่ recalc แล้ว:", rc)
    else:
        shutil.rmtree(tmp, ignore_errors=True)
    print("ผล:", "ผ่าน" if not errs and not mism else "ไม่ผ่าน")
    return 1 if errs or mism else 0


if __name__ == "__main__":
    sys.exit(main())
