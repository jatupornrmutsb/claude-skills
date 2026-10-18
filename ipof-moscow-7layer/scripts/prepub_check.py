#!/usr/bin/env python3
"""สแกนไฟล์ .docx ก่อนส่งทุกครั้ง ตามกฎการเขียนของ Skill ipof-moscow-7layer

  python prepub_check.py บทความ_V3.docx [ไฟล์อื่น.docx ...] [--excel สรุปผล_V2.xlsx ...]
        [--forbid "คำ1|คำ2"] [--allow-colon "คำสำคัญ|Keywords"] [--render OUTDIR]

ตรวจ
- คำที่ควรหายไปแล้ว (--forbid)
- ":" ต่อจากภาษาไทย (ยกเว้นป้ายตามแม่แบบ)
- อ้างอิงแบบช่วง [1]–[3] / [1-3] หรือหลายแหล่งในวงเล็บเดียว [1, 2]
- เลขอ้างอิงเรียงตามการอ้างครั้งแรก, อ้างแล้วไม่มีรายการ / มีรายการไม่ถูกอ้าง
- รายการอ้างอิงที่เป็นเว็บเพจ/บล็อก (ไม่มี DOI/ชื่อวารสาร/การประชุม)
- เลข "ตารางที่ n" และ "รูปที่ n" เรียง 1..n
- รูปแบบ “ระบบต้อง…” ในเครื่องหมายคำพูด
- ถ้อยคำเกินหลักฐาน (ติดตั้งใช้งานแล้ว ฯลฯ) — ต้องคงระดับแนวคิด/ต้นแบบ
- ช่วงตัวเลขที่ใช้ "-" แทน "–"
- ตัวเลขทศนิยม 2 ตำแหน่งที่ไม่พบในไฟล์ Excel (หลัง recalc)
- --project: ชื่อเรื่องในไฟล์ต้องตรงกับ project.title ที่ยืนยันแล้ว (ไม่ถูกเปลี่ยน)
- --render: แปลงเป็น PDF รายงานจำนวนหน้า
ออกด้วยรหัส 1 เมื่อมี error
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import os
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
import tempfile
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))

TH = r"[฀-๿]"
REF_HEAD = re.compile(r"^\s*(\d+\.?\s*)?(เอกสารอ้างอิง|References|REFERENCES)\s*$")
OVERCLAIM = ["ติดตั้งใช้งานแล้ว", "ใช้งานจริงแล้ว", "นำไปใช้งานจริงแล้ว", "เปิดใช้งานแล้ว", "ระบบที่ใช้งานอยู่", "deployed", "in production"]


def paragraphs(doc):
    out = []
    for el in doc.element.body.iter(qn("w:p")):
        t = []
        for n in el.iter():
            if n.tag == qn("w:t"):
                t.append(n.text or "")
            elif n.tag in (qn("w:br"), qn("w:cr")):
                t.append("\n")
        out.append("".join(t))
    return out


def norm_title(t):
    return re.sub(r"\s+", " ", (t or "").replace("\n", " ")).strip()


def check_title(paras, title, E, W):
    want = norm_title(title)
    if not want:
        W("TITLE_UNSET", "ยังไม่มีชื่อเรื่องใน project.json — ยืนยันชื่อเรื่องกับผู้ใช้ก่อน")
        return
    texts = [norm_title(t) for t in paras]
    if want in texts or any(want in t for t in texts):
        return
    first = next((t for t in texts if t), "")
    near = max(texts, key=lambda t: len(set(t.split()) & set(want.split())) if t else -1) if texts else ""
    E("TITLE_MISMATCH", f"ไม่พบชื่อเรื่องที่ยืนยันไว้ “{want[:60]}” — ย่อหน้าแรก “{first[:60]}” / ใกล้เคียงที่สุด “{near[:60]}”")


def check_file(path, forbid, allow_colon, issues, excel_vals, title=None):
    doc = Document(path)
    paras = paragraphs(doc)
    name = Path(path).name
    E = lambda code, msg: issues.append(("error", name, code, msg))  # noqa: E731
    W = lambda code, msg: issues.append(("warn", name, code, msg))  # noqa: E731
    ref_i = next((i for i, t in enumerate(paras) if REF_HEAD.match(t.strip())), None)
    body = paras if ref_i is None else paras[:ref_i]
    refs = [] if ref_i is None else [t.strip() for t in paras[ref_i + 1:] if re.match(r"^\s*\[\d+\]", t)]
    full = "\n".join(paras)
    if title is not False:
        check_title(paras, title, E, W)
    if forbid:
        for w in forbid.split("|"):
            n = full.count(w)
            if w and n:
                E("FORBIDDEN", f"ยังพบ “{w}” {n} จุด")
    allow = re.compile(rf"({allow_colon})\s*[:：]") if allow_colon else None
    for i, t in enumerate(paras):
        for m in re.finditer(rf"{TH}[ะ-๎]*\s*[:：]", t):
            ctx = t[max(0, m.start() - 15): m.end() + 5]
            if allow and allow.search(t[max(0, m.start() - 20): m.end()]):
                continue
            E("THAI_COLON", f"“:” ต่อจากภาษาไทย: …{ctx}…")
        for m in re.finditer(r"“[^”]{0,6}ระบบต้อง[^”]*”|\"[^\"]{0,6}ระบบต้อง[^\"]*\"", t):
            E("QUOTED_SHALL", f"ใช้ “ระบบต้อง…” ในเครื่องหมายคำพูด: {m.group(0)[:50]}")
        for w in OVERCLAIM:
            if w in t:
                E("OVERCLAIM", f"ถ้อยคำเกินหลักฐาน “{w}”: {t[:60]}")
        for m in re.finditer(r"\b\d+\.\d{2}\s*-\s*\d+\.\d{2}\b", t):
            W("DASH", f"ช่วงตัวเลขใช้ - ควรเป็น –: {m.group(0)}")
    # อ้างอิง
    cites = []
    for t in body:
        for m in re.finditer(r"\[(\d{1,3})\]\s*[–\-–—]\s*\[(\d{1,3})\]", t):
            E("CITE_RANGE", f"อ้างแบบช่วง {m.group(0)} — ต้องอ้างแหล่งเดียวต่อข้อความ")
        for m in re.finditer(r"\[(\d{1,3}(?:\s*[,–\-–]\s*\d{1,3})+)\]", t):
            E("CITE_MULTI", f"หลายแหล่งในวงเล็บเดียว {m.group(0)}")
        for m in re.finditer(r"\[(\d{1,3})\]", t):
            cites.append(int(m.group(1)))
    first = list(OrderedDict.fromkeys(cites))
    if first and first != list(range(1, len(first) + 1)):
        E("CITE_ORDER", f"ลำดับอ้างครั้งแรกไม่เรียง: {first[:20]}")
    listed = [int(re.match(r"^\s*\[(\d+)\]", r).group(1)) for r in refs]
    if refs:
        if listed != list(range(1, len(listed) + 1)):
            E("REF_NUMBERING", f"เลขรายการไม่เรียง: {listed}")
        miss = sorted(set(first) - set(listed))
        unused = sorted(set(listed) - set(first))
        if miss:
            E("REF_MISSING", f"อ้างแต่ไม่มีในรายการ: {miss}")
        if unused:
            E("REF_UNUSED", f"มีในรายการแต่ไม่ถูกอ้าง: {unused}")
        for r in refs:
            pub = re.search(r"doi|vol\.|no\.|pp\.|Proc\.|Conf|Journal|J\.|Trans\.|ISBN|วารสาร|การประชุม|รายงานการวิจัย|วิทยานิพนธ์|สารนิพนธ์|thesis|Press|Publisher|สำนักพิมพ์", r, re.I)
            web = re.search(r"https?://|www\.|เว็บไซต์|blog|บล็อก|\[Online\]|Accessed|สืบค้นเมื่อ", r, re.I)
            if web and not pub:
                E("REF_WEB", f"รายการเป็นเว็บเพจ/ไม่พบข้อมูลการตีพิมพ์: {r[:70]}")
    elif ref_i is None and cites:
        W("REF_HEADING", "ไม่พบหัวข้อเอกสารอ้างอิง")
    for label in ("ตารางที่", "รูปที่"):
        nums = []
        for t in body:
            m = re.match(rf"^\s*{label}\s*(\d+)", t)
            if m:
                nums.append(int(m.group(1)))
        if nums and nums != list(range(1, len(nums) + 1)):
            E("CAPTION_ORDER", f"เลข{label}ไม่เรียง: {nums}")
        refd = sorted({int(x) for t in body for x in re.findall(rf"{label}\s*(\d+)", t)})
        if nums and set(refd) - set(nums):
            E("CAPTION_REF", f"อ้าง{label} {sorted(set(refd) - set(nums))} ที่ไม่มีชื่อ{label}")
    if excel_vals is not None:
        nums = set()
        for t in paras:
            for m in re.finditer(r"(?<![\d.])(\d{1,3}\.\d{2})(?![\d.])", t):
                nums.add(m.group(1))
        missing = sorted(n for n in nums if n not in excel_vals)
        if missing:
            W("NUM_NOT_IN_EXCEL", f"ตัวเลขที่ไม่พบในไฟล์ Excel ({len(missing)}): {missing[:25]} — ตรวจว่าไม่ใช่ผลคำนวณ หรือแก้ให้ตรง")
    return len(paras)


def excel_values(paths):
    from verify_excel import recalc
    from openpyxl import load_workbook
    vals = set()
    tmp = Path(tempfile.mkdtemp(prefix="xl_"))
    for p in paths:
        rc = recalc(p, tmp)
        wb = load_workbook(rc, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        vals.add(f"{v:.2f}")
                    elif isinstance(v, str):
                        vals.update(re.findall(r"\d+\.\d{2}", v))
    shutil.rmtree(tmp, ignore_errors=True)
    return vals


def render_pages(path, outdir):
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    prof = tempfile.mkdtemp(prefix="lo_")
    subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)],
                   capture_output=True, timeout=240)
    pdf = Path(outdir) / (Path(path).stem + ".pdf")
    if not pdf.exists():
        return None
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONUTF8": "1"}).stdout
    m = re.search(r"Pages:\s+(\d+)", info)
    return (int(m.group(1)) if m else None), pdf


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx", nargs="+")
    ap.add_argument("--excel", nargs="*")
    ap.add_argument("--forbid", default="")
    ap.add_argument("--allow-colon", default="คำสำคัญ|Keywords|บทคัดย่อ|Abstract|ผู้ประพันธ์บรรณกิจ|E-mail|อีเมล")
    ap.add_argument("--render", metavar="OUTDIR")
    ap.add_argument("--project", help="project.json — ตรวจชื่อเรื่องไม่ถูกเปลี่ยน")
    ap.add_argument("--no-title-check", nargs="*", default=[], metavar="FILE",
                    help="ไฟล์ที่ไม่ต้องมีชื่อเรื่อง (เช่น ตรวจเทียบอ้างอิง)")
    a = ap.parse_args(argv)
    ev = excel_values(a.excel) if a.excel else None
    issues = []
    title = False
    if a.project:
        import json
        title = json.loads(Path(a.project).read_text(encoding="utf-8")).get("project", {}).get("title")
    skip = {Path(x).name for x in a.no_title_check}
    for f in a.docx:
        check_file(f, a.forbid, a.allow_colon, issues, ev, False if Path(f).name in skip else title)
        if a.render:
            Path(a.render).mkdir(parents=True, exist_ok=True)
            r = render_pages(f, a.render)
            print(f"{Path(f).name}: " + (f"{r[0]} หน้า (PDF: {r[1]})" if r else "เรนเดอร์ไม่ได้ — ยังไม่ได้ตรวจหน้าจริง"))
    for lvl, name, code, msg in issues:
        print(f"[{lvl}] {name} {code}: {msg}")
    ne = sum(1 for i in issues if i[0] == "error")
    print(f"สรุป: error {ne}, warn {len(issues) - ne}")
    return 1 if ne else 0


if __name__ == "__main__":
    sys.exit(main())
