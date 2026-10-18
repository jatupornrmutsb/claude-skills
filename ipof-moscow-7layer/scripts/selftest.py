#!/usr/bin/env python3
"""ทดสอบสคริปต์ทั้งหมดด้วยแบบสอบถาม/แบบประเมินสมมุติที่สร้างขึ้นเอง (ไม่ใช้ข้อมูลจริง)

  python selftest.py [--workdir DIR] [--no-render]
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import shutil
import statistics as S
import subprocess
import os
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
EX = HERE.parent / "assets" / "examples" / "project.example.json"
PY = sys.executable
N_TEST_USERS = 12  # จำนวนผู้ตอบสมมติสำหรับทดสอบสคริปต์เท่านั้น
RES = []


def check(name, cond, detail=""):
    RES.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + ("" if cond or not detail else f" — {str(detail)[-500:]}"))


def sh(*a):
    r = subprocess.run([PY, *map(str, a)], capture_output=True, text=True, encoding="utf-8", errors="replace", env={**os.environ, "PYTHONUTF8": "1"})
    return r.returncode, r.stdout + r.stderr


def add_sym_box(par, checked):
    r = par.add_run()
    sym = OxmlElement("w:sym")
    sym.set(qn("w:font"), "Wingdings")
    sym.set(qn("w:char"), "F0FE" if checked else "F0A8")
    r._r.append(sym)


def box(par, checked, style):
    if style == "sym":
        add_sym_box(par, checked)
    else:
        par.add_run("☑" if checked else "☐")


def score_table(doc, items, answers, style, levels=(5, 4, 3, 2, 1)):
    t = doc.add_table(rows=1, cols=2 + len(levels))
    t.style = "Table Grid"
    hdr = ["รหัส", "รายการ"] + [str(x) for x in levels]
    for j, h in enumerate(hdr):
        t.rows[0].cells[j].text = h
    for code, text in items:
        row = t.add_row().cells
        row[0].text = code
        row[1].text = text
        v = answers.get(code)
        for j, lv in enumerate(levels, 2):
            p = row[j].paragraphs[0]
            chosen = (lv in v) if isinstance(v, (list, tuple)) else (v == lv)
            box(p, chosen, style)


def confirm_table(doc, items, ans, style, with_level=True):
    cols = 7 + 3 if with_level else 4
    t = doc.add_table(rows=2 if with_level else 1, cols=(10 if with_level else 5))
    t.style = "Table Grid"
    if with_level:
        h0 = ["รหัส", "ข้อกำหนด", "ระดับเดิม", "เห็นชอบ", "ไม่เห็นชอบ", "ระดับที่เสนอ", "", "", "", "เหตุผล"]
        for j, h in enumerate(h0):
            t.rows[0].cells[j].text = h
        t.rows[0].cells[5].merge(t.rows[0].cells[8])
        t.rows[0].cells[5].text = "ระดับที่เสนอ"
        for j, h in enumerate(["", "", "", "", "", "M", "S", "C", "W", ""]):
            t.rows[1].cells[j].text = h
    else:
        for j, h in enumerate(["รหัส", "รายการ", "เห็นชอบ", "ไม่เห็นชอบ", "เหตุผล"]):
            t.rows[0].cells[j].text = h
    for code, text, lvl in items:
        a = ans[code]
        row = t.add_row().cells
        row[0].text = code
        row[1].text = text
        off = 0
        if with_level:
            row[2].text = lvl
            off = 1
        box(row[2 + off].paragraphs[0], a["agree"], style)
        box(row[3 + off].paragraphs[0], not a["agree"], style)
        if with_level:
            for j, lv in enumerate("MSCW", 5):
                box(row[j].paragraphs[0], a.get("prop") == lv, style)
            row[9].text = a.get("reason", "")
        else:
            row[4].text = a.get("reason", "")


def choice_line(doc, question, options, chosen, style):
    doc.add_paragraph(question)
    for o in options:
        p = doc.add_paragraph()
        box(p, o in chosen, style)
        p.add_run(" " + o)


def make_user_forms(proj, d, rng):
    truth = {}
    funcs = [f["text"] for f in proj["functions"]]
    roles = ["ผู้ดูแลเครือข่าย", "ผู้บริหาร", "เจ้าหน้าที่สนับสนุน"]
    items = [(x["code"], x["text"]) for x in proj["problems"] + proj["needs"]]
    choices_truth = {}
    for i in range(1, N_TEST_USERS + 1):
        rid = f"U{i:02d}"
        style = "sym" if i % 3 == 0 else "text"
        ans = {c: rng.choice([3, 4, 4, 5, 5]) for c, _ in items}
        if i == 4:
            ans["N03"] = None  # ไม่ติ๊ก
        if i == 5:
            ans["N04"] = [4, 5]  # ติ๊กเกิน
        truth[rid] = {k: v for k, v in ans.items() if isinstance(v, int)}
        doc = Document()
        doc.add_paragraph("แบบสอบถามความต้องการ (ตัวอย่างสมมุติ)")
        role = rng.choice(roles)
        choice_line(doc, "1. บทบาท", roles, [role], style)
        picks = [f for f in funcs if rng.random() < 0.6]
        choices_truth[rid] = picks
        score_table(doc, items, ans, style)
        choice_line(doc, "ฟังก์ชันที่ควรพัฒนาก่อน (เลือกได้หลายข้อ)", funcs, picks, style)
        doc.add_paragraph("OP01 ปัญหาอื่นที่พบในการใช้ข้อมูลสถานะอุปกรณ์")
        doc.add_paragraph(f"อุปกรณ์รุ่นเก่าส่งข้อมูลไม่ได้ (ผู้ตอบ {rid})" if i % 2 else "........")
        doc.save(d / f"{rid}.docx")
    return truth, choices_truth


def make_expert1(proj, d, rng):
    r1 = proj["round1"]
    text = {r["code"]: r["text"] for r in proj["requirements"]}
    truth = {}
    plan = {"FR06": [True, True, True, False, False],  # 3/5 → ปรับตามเสียงข้างมาก
            "NFR01": [True, True, True, True, False]}  # 4/5 → ถึงเกณฑ์
    for e in range(1, 6):
        eid = f"E{e:02d}"
        style = "sym" if e == 2 else "text"
        doc = Document()
        choice_line(doc, "ตำแหน่ง", ["หัวหน้างาน", "นักวิชาการคอมพิวเตอร์", "อาจารย์"], ["นักวิชาการคอมพิวเตอร์" if e % 2 else "อาจารย์"], style)
        sc = {}
        for c in r1["score_items"]:
            v = rng.choice([4, 5])
            sc[c] = v
            doc.add_paragraph(f"{c} {text[c]}")
            p = doc.add_paragraph("คะแนนความเหมาะสม ")
            for lv in (5, 4, 3, 2, 1):
                box(p, lv == v, style)
                p.add_run(f" {lv} ")
        ans = {}
        for c in r1["confirm_items"]:
            ag = plan.get(c, [True] * 5)[e - 1]
            ans[c] = {"agree": ag, "prop": None if ag else ("M" if c == "FR06" else "S"), "reason": "" if ag else "ค้นหาเป็นงานหลักของผู้ดูแล"}
        confirm_table(doc, [(c, text[c], r1["levels"][c]) for c in r1["confirm_items"]], ans, style)
        choice_line(doc, "ความเห็นโดยรวม", ["เหมาะสม", "เหมาะสมแต่ควรปรับปรุง", "ไม่เหมาะสม"], ["เหมาะสม"], style)
        doc.add_paragraph("OP01 ข้อเสนอแนะเพิ่มเติม")
        doc.add_paragraph("ควรระบุแหล่งข้อมูลต้นทางของแต่ละศูนย์")
        doc.save(d / f"{eid}.docx")
        truth[eid] = {"scores": sc, "agree": {c: a["agree"] for c, a in ans.items()}}
    return truth


def make_expert2(proj, d, rng):
    truth = {}
    for e in range(1, 6):
        eid = f"E{e:02d}"
        doc = Document()
        a = {x["code"]: rng.choice([4, 5]) for x in proj["arch_items"] + proj["eval_items"]}
        doc.add_paragraph("ส่วนที่ 1 สถาปัตยกรรม")
        score_table(doc, [(x["code"], x["text"]) for x in proj["arch_items"]], a, "text")
        doc.add_paragraph("ส่วนที่ 2 ข้อกำหนด")
        score_table(doc, [(x["code"], x["text"]) for x in proj["eval_items"]], a, "text")
        ans = {c["code"]: {"agree": True} for c in proj["confirm_items"]}
        confirm_table(doc, [(c["code"], c["text"], "") for c in proj["confirm_items"]], ans, "text", with_level=False)
        choice_line(doc, "ความเห็นโดยรวม", ["เหมาะสม", "เหมาะสมแต่ควรปรับปรุง", "ไม่เหมาะสม"], ["เหมาะสม"], "text")
        doc.save(d / f"{eid}.docx")
        truth[eid] = a
    return truth


def load_csv(p):
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workdir")
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    wd = Path(a.workdir or tempfile.mkdtemp(prefix="ipof_selftest_"))
    if wd.exists():
        shutil.rmtree(wd)
    wd.mkdir(parents=True)
    proj = json.loads(EX.read_text(encoding="utf-8"))
    rng = random.Random(7)

    # 0) ทะเบียนตัวอย่างผ่านการตรวจ
    rc, out = sh(HERE / "analyze.py", "validate", "--project", EX)
    check("validate: ทะเบียนตัวอย่างไม่มี error", rc == 0, out)
    check("validate: เตือนเมื่อยังไม่ได้ถามคำถามหลัก/n/k/วารสาร", "จำนวนผู้ตอบ (n)" in out and "วารสารเป้าหมาย" in out, out)
    bad = json.loads(json.dumps(proj))
    bad["requirements"][2]["component"] = ""
    bad["requirements"][3]["depends_on"] = ["FR02"]  # Must พึ่งพา Should
    bad["requirements"][6]["moscow_reason"] = ""
    bad["requirements"][9]["group"] = "G1"
    (wd / "bad.json").write_text(json.dumps(bad, ensure_ascii=False), encoding="utf-8")
    rc, out = sh(HERE / "analyze.py", "validate", "--project", wd / "bad.json")
    check("validate: จับ Must ไม่มีองค์ประกอบ / Must พึ่งพา Should / ระดับต่างไม่มีเหตุผล / NFR มีกลุ่ม",
          rc == 1 and "องค์ประกอบ" in out and "พึ่งพา FR02" in out and "moscow_reason" in out and "NFR ไม่จัดกลุ่ม" in out, out)

    # 1) แบบสอบถามผู้ใช้
    fu = wd / "forms_user"
    fu.mkdir()
    truth, ch_truth = make_user_forms(proj, fu, rng)
    du = wd / "data_user"
    rc, out = sh(HERE / "read_forms.py", du, fu)
    check("read_forms: รายงานแถวไม่ติ๊ก/ติ๊กเกิน (exit 1)", rc == 1 and "ไม่ได้ติ๊ก" in out and "ติ๊ก 2 ช่อง" in out, out)
    got = {}
    for r in load_csv(du / "scores.csv"):
        got.setdefault(r["respondent"], {})[r["item"]] = int(r["value"])
    check("read_forms: คะแนนทุกค่าตรงกับที่สร้าง (☑ และ Wingdings)", got == truth,
          [(k, set(truth[k].items()) ^ set(got.get(k, {}).items())) for k in truth if truth[k] != got.get(k)][:3])
    chs = load_csv(du / "choices.csv")
    fsel = {}
    for c in chs:
        if "ฟังก์ชัน" in c["question"] and c["checked"] == "1":
            fsel.setdefault(c["respondent"], []).append(c["option"])
    check("read_forms: คำถามเลือกหลายข้อถูกต้อง", all(sorted(fsel.get(k, [])) == sorted(v) for k, v in ch_truth.items()))
    txt = load_csv(du / "text.csv")
    check("read_forms: คำตอบปลายเปิด OP01 (ข้ามบรรทัดจุด)", len([t for t in txt if t["item"] == "OP01"]) == 6, txt[:3])

    # 2) สถิติ: pooled จากคะแนนดิบ, S.D. n−1
    rc, out = sh(HERE / "analyze.py", "user", "--project", EX, "--data", du, "--json", wd / "res_user.json")
    res = json.loads((wd / "res_user.json").read_text(encoding="utf-8"))
    n01 = [truth[r]["N01"] for r in truth]
    check("สถิติรายข้อ: X̄/S.D.(n−1) ตรง statistics", abs(res["items"]["N01"]["mean"] - S.mean(n01)) < 1e-9 and abs(res["items"]["N01"]["sd"] - S.stdev(n01)) < 1e-9)
    d1 = [n["code"] for n in proj["needs"] if n["dimension"] == "D1"]
    raw = [truth[r][c] for r in truth for c in d1 if c in truth[r]]
    mean_of_means = S.mean([S.mean([truth[r][c] for r in truth if c in truth[r]]) for c in d1])
    check("รายด้านใช้คะแนนดิบทุกค่า (ไม่ใช่เฉลี่ยของค่าเฉลี่ย)", abs(res["dimensions"]["ด้านข้อมูล"]["mean"] - S.mean(raw)) < 1e-9 and
          abs(S.mean(raw) - mean_of_means) > 1e-6, (S.mean(raw), mean_of_means))
    fr01 = [truth[r][c] for r in truth for c in ("N01", "N02") if c in truth[r]]
    check("ค่าเฉลี่ยข้อกำหนดรวมหลาย N = pooled คะแนนดิบ", abs(res["requirements"]["FR01"]["mean"] - S.mean(fr01)) < 1e-9)
    check("MoSCoW: รายงานข้อที่ระดับต่างจากคะแนนพร้อมเหตุผล", "ฉบับสุดท้าย" in out, out[-600:])

    # 3) Excel ผู้ใช้ — สูตร + ตรวจแยก
    xu = wd / "สรุปผลผู้ใช้_V1.xlsx"
    rc, out = sh(HERE / "build_excel.py", "--mode", "user", "--project", EX, "--data", du, "--out", xu)
    check("build_excel user", rc == 0 and xu.exists(), out)
    rc, out = sh(HERE / "verify_excel.py", xu)
    check("verify_excel user: ไม่มี error และตัวเลขตรงการคำนวณแยก", rc == 0, out)
    from openpyxl import load_workbook
    wb = load_workbook(xu)
    names = wb.sheetnames
    check("Excel ผู้ใช้มีแผ่นครบ", all(s in names for s in ["สรุปผล", "ข้อมูลทั่วไป", "ข้อมูลดิบ", "ค่าเฉลี่ยรายด้าน", "การเลือกฟังก์ชัน", "คำตอบปลายเปิด"]), names)
    ws = wb["ค่าเฉลี่ยรายด้าน"]
    check("Excel: ค่าเป็นสูตร ฟอนต์ TH Sarabun New ข้อมูลดิบสีน้ำเงิน",
          str(ws["D4"].value).startswith("=AVERAGE") and ws["D4"].font.name == "TH Sarabun New" and wb["ข้อมูลดิบ"]["B4"].font.color.rgb.endswith("0000FF"))

    # 4) ผู้เชี่ยวชาญครั้งที่ 1
    f1 = wd / "forms_e1"
    f1.mkdir()
    t1 = make_expert1(proj, f1, rng)
    d1d = wd / "data_e1"
    rc, out = sh(HERE / "read_forms.py", d1d, f1)
    check("read_forms ครั้งที่ 1: บรรทัด 'คะแนนความเหมาะสม' + ตารางหัว 2 ชั้น", rc == 0, out)
    sc1 = {}
    for r in load_csv(d1d / "scores.csv"):
        sc1.setdefault(r["respondent"], {})[r["item"]] = r["value"]
    ok = all(int(sc1[e][c]) == v for e in t1 for c, v in t1[e]["scores"].items()) and \
        all(sc1[e][f"{c}#agree"] == ("agree" if v else "disagree") for e in t1 for c, v in t1[e]["agree"].items())
    check("ครั้งที่ 1: คะแนนและการเห็นชอบตรงทุกค่า", ok)
    check("ครั้งที่ 1: อ่านระดับที่เสนอ (M) เฉพาะผู้ไม่เห็นชอบ", sum(1 for e in sc1 if sc1[e].get("FR06#moscow") == "M") == 2)
    rc, out = sh(HERE / "analyze.py", "agree", "--data", d1d, "--items", "FR01,FR06")
    check("เกณฑ์ 80%: FR06 = ปรับตามเสียงข้างมาก (3/5), FR01 ถึงเกณฑ์", "ปรับตามเสียงข้างมาก (3/5)" in out and "FR01: เห็นชอบ 5/5" in out, out)
    x1 = wd / "สรุปผลผู้เชี่ยวชาญครั้งที่1_V1.xlsx"
    rc, out = sh(HERE / "build_excel.py", "--mode", "expert1", "--project", EX, "--data", d1d, "--out", x1)
    rc2, out2 = sh(HERE / "verify_excel.py", x1, "--keep-recalc")
    check("Excel ครั้งที่ 1: ไม่มี error และตัวเลขตรง", rc == 0 and rc2 == 0, out + out2)
    rcp = out2.split("สำเนาที่ recalc แล้ว:")[-1].strip().splitlines()[0] if "สำเนาที่ recalc" in out2 else None
    if rcp:
        w = load_workbook(rcp, data_only=True)["ยืนยัน MoSCoW"]
        rows = {w.cell(r, 1).value: [w.cell(r, c).value for c in range(1, 14)] for r in range(4, 11)}
        check("Excel ครั้งที่ 1: FR06 ผล/ระดับฉบับสุดท้ายถูก", rows["FR06"][11] == "ปรับตามเสียงข้างมาก (3/5)" and rows["FR06"][12] == "M", rows.get("FR06"))
        check("Excel ครั้งที่ 1: NFR01 ถึงเกณฑ์ (4/5)", rows["NFR01"][11] == "ถึงเกณฑ์ (4/5)", rows.get("NFR01"))
    check("Excel ครั้งที่ 1 แผ่นครบ", all(s in load_workbook(x1).sheetnames for s in
          ["สรุปผล", "ข้อมูลผู้เชี่ยวชาญ", "ข้อมูลดิบ", "ค่าเฉลี่ย", "ยืนยัน MoSCoW", "ความเห็นโดยรวม", "เหตุผลรายข้อ", "คำตอบปลายเปิด", "สังเคราะห์ประเด็น"]))

    # 5) ผู้เชี่ยวชาญครั้งที่ 2 + เทียบชุดใหม่
    f2 = wd / "forms_e2"
    f2.mkdir()
    make_expert2(proj, f2, rng)
    d2 = wd / "data_e2"
    rc, out = sh(HERE / "read_forms.py", d2, f2)
    x2 = wd / "สรุปผลผู้เชี่ยวชาญครั้งที่2_V1.xlsx"
    rcb, outb = sh(HERE / "build_excel.py", "--mode", "expert2", "--project", EX, "--data", d2, "--out", x2)
    rcv, outv = sh(HERE / "verify_excel.py", x2)
    check("Excel ครั้งที่ 2: สร้างและตรวจผ่าน", rc == 0 and rcb == 0 and rcv == 0, out + outb + outv)
    check("Excel ครั้งที่ 2 แผ่นครบ", all(s in load_workbook(x2).sheetnames for s in
          ["สรุปผล", "ข้อมูลผู้เชี่ยวชาญ", "ข้อมูลดิบ-คะแนน", "ค่าเฉลี่ย-สถาปัตยกรรม", "ค่าเฉลี่ย-ข้อกำหนด", "การยืนยัน", "ความเห็นโดยรวม", "ข้อเสนอแนะ", "สังเคราะห์ประเด็น"]))
    d2b = wd / "data_e2_new"
    shutil.copytree(d2, d2b)
    rows = load_csv(d2b / "scores.csv")
    rows[0]["value"] = "3" if rows[0]["value"] != "3" else "4"
    with open(d2b / "scores.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["respondent", "item", "value"])
        w.writeheader()
        w.writerows(rows)
    rc, out = sh(HERE / "analyze.py", "compare", d2, d2b)
    check("compare: รายงานข้อที่เปลี่ยนในชุดใหม่", "เปลี่ยน 1 ค่า" in out and rows[0]["item"] in out, out)

    # 6) แผนภาพ
    if not a.no_render:
        dg = wd / "diagrams"
        rc, out = sh(HERE / "render_diagrams.py", "--project", EX, "--outdir", dg)
        pngs = sorted(p.name for p in dg.glob("*.png"))
        check("แผนภาพ: ipof/layers7/context/dataflow/traceability เป็น PNG และไม่มีปัญหาความครบถ้วน",
              rc == 0 and {"ipof.png", "layers7.png", "context.png", "dataflow.png", "traceability.png"} <= set(pngs), out)
        h = (dg / "layers7.html").read_text(encoding="utf-8")
        check("ภาพ 7 ชั้น: ข้อกำหนดทุกข้ออยู่ในภาพพร้อมระดับ", all(f"<b>{r['code']}</b> [{r['moscow_final']}]" in h for r in proj["requirements"]))
        ip = (dg / "ipof.html").read_text(encoding="utf-8")
        check("IPOF: มีแถบกำกับ + คำอธิบายระดับฉบับสุดท้าย + Output เฉพาะ Must G3/G4",
              "ครอบทุกองค์ประกอบ" in ip and "ระดับฉบับสุดท้าย" in ip and "<b>FR06</b>" not in ip and "<b>FR04</b>" in ip)
        rc, out = sh(HERE / "render_diagrams.py", "--project", wd / "bad.json", "--outdir", wd / "dg_bad", "--no-png", "--only", "layers")
        check("แผนภาพ: จับ Must ที่ไม่มีองค์ประกอบ", rc == 1 and "FR03" in out, out)

    # 7) แก้ docx คงรูปแบบ + รุ่น
    art = wd / "บทความ.docx"
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("ผลการประเมินพบว่า ค่าเฉลี่ย")
    r1.bold = True
    p.add_run("ภาพรวม 4.")
    p.add_run("21")
    br_run = p.add_run()
    br_run._r.append(OxmlElement("w:br"))
    p.add_run("อยู่ในระดับมากที่สุด")
    doc.add_paragraph("ตารางที่ 1 ปัญหาและความต้องการ")
    t = doc.add_table(rows=2, cols=2)
    t.cell(1, 1).text = "4.00"
    doc.save(art)
    edits = wd / "edits.json"
    edits.write_text(json.dumps([{"find": "4.21", "replace": "4.25", "count": 1}, {"table": 1, "row": 2, "col": 2, "value": "4.10"}], ensure_ascii=False), encoding="utf-8")
    rc, out = sh(HERE / "docx_edit.py", art, edits)
    v1 = wd / "บทความ_V1.docx"
    d = Document(v1) if v1.exists() else None
    ok = d is not None and "4.25" in d.paragraphs[0].text and d.paragraphs[0].runs[0].bold and \
        d.paragraphs[0]._p.find(".//" + qn("w:br")) is not None and d.tables[0].cell(1, 1).text == "4.10"
    check("docx_edit: แทนข้าม run, คงตัวหนาและการขึ้นบรรทัด, แก้เซลล์, ตั้งชื่อ _V1", rc == 0 and ok, out)
    rc, out = sh(HERE / "docx_edit.py", v1, edits)
    check("docx_edit: count ไม่ตรง → ไม่บันทึก", rc == 1 and not (wd / "บทความ_V2.docx").exists(), out)
    edits.write_text(json.dumps([{"find": "มากที่สุด", "replace": "มากที่สุด"}], ensure_ascii=False), encoding="utf-8")
    sh(HERE / "docx_edit.py", v1, edits)
    check("docx_edit: รุ่นถัดไป _V1 → _V2", (wd / "บทความ_V2.docx").exists())

    # 8) สแกนก่อนส่ง
    bad_doc = wd / "scan_bad.docx"
    doc = Document()
    for s in ["คำสำคัญ: แดชบอร์ด, MoSCoW",
              "ผลการวิจัย: พบว่าระบบมีความเหมาะสม [2]",
              "งานก่อนหน้า [1]–[3] และ [4, 5] ระบุว่า “ระบบต้องแจ้งเตือน”",
              "ระบบติดตั้งใช้งานแล้วในทุกศูนย์ ช่วงคะแนน 3.41-4.20 ค่าเฉลี่ย 4.97",
              "ตารางที่ 2 ผลการประเมิน",
              "เอกสารอ้างอิง",
              "[1] A. Author, “Paper,” IEEE Access, vol. 1, pp. 1–2, 2020.",
              "[2] Blog, “How to MoSCoW,” Available: https://example.com/blog (accessed Jan. 1, 2024).",
              "[3] C. Author, Book. Publisher, 2019."]:
        doc.add_paragraph(s)
    doc.save(bad_doc)
    rc, out = sh(HERE / "prepub_check.py", bad_doc, "--forbid", "ทุกศูนย์", "--excel", xu)
    for code in ["FORBIDDEN", "THAI_COLON", "CITE_RANGE", "CITE_MULTI", "CITE_ORDER", "QUOTED_SHALL", "OVERCLAIM", "DASH", "REF_WEB", "CAPTION_ORDER", "NUM_NOT_IN_EXCEL"]:
        check(f"prepub_check จับ {code}", code in out, out)
    check("prepub_check: ไม่จับ 'คำสำคัญ:' ตามแม่แบบ", "คำสำคัญ: แดช" not in out)
    good = wd / "scan_good.docx"
    doc = Document()
    for s in ["คำสำคัญ: แดชบอร์ด", "ระบบต้นแบบในระดับแนวคิดได้รับการประเมิน [1] และใช้กรอบ MoSCoW [2]", "ตารางที่ 1 ผล",
              "เอกสารอ้างอิง", "[1] A. Author, “P,” J. Test, vol. 1, 2020.", "[2] B. Author, “Q,” in Proc. Conf., 2021, pp. 1–2."]:
        doc.add_paragraph(s)
    doc.save(good)
    rc, out = sh(HERE / "prepub_check.py", good, "--render", wd / "pdf")
    check("prepub_check: เอกสารที่ถูกกฎผ่าน + รายงานจำนวนหน้า", rc == 0 and "หน้า" in out, out)

    n = sum(1 for _, c in RES if c)
    print(f"\nสรุป: ผ่าน {n}/{len(RES)} (workdir: {wd})")
    return 0 if n == len(RES) else 1


if __name__ == "__main__":
    sys.exit(main())
