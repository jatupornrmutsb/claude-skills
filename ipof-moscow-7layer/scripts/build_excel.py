#!/usr/bin/env python3
"""สร้างไฟล์ Excel สรุปผล (ทุกค่าเป็นสูตร) จากผลของ read_forms.py + project.json

  python build_excel.py --mode user|expert1|expert2 --project project.json \
      --data DIR_จาก_read_forms --out สรุปผล_V1.xlsx

กติกา
- ข้อมูลนำเข้า (คะแนนดิบ ช่องเลือก ข้อความ) = ตัวอักษรสีน้ำเงิน, สูตร = สีดำ
- X̄ = AVERAGE, S.D. = STDEV.S (n − 1); รายด้าน/ภาพรวม = AVERAGE/STDEV.S ของช่วงคะแนนดิบทุกข้อในด้านนั้น
  (ไม่เฉลี่ยค่าเฉลี่ยรายข้อ)
- แปลผลด้วยสูตร IF จากค่าที่ปัด 2 ตำแหน่ง
- แถวรวม/รายด้านมีพื้นสี, ฟอนต์ TH Sarabun New
- เขียน <out>.manifest.json สำหรับ verify_excel.py ตรวจตัวเลขแยกอีกทาง
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys

for _s in (sys.stdout, sys.stderr):  # Windows console ไม่ใช่ UTF-8 จะ print ภาษาไทยไม่ได้
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ipof_lib as L  # noqa: E402

FONT = "TH Sarabun New"
F_IN = Font(name=FONT, size=14, color="0000FF")
F_N = Font(name=FONT, size=14)
F_B = Font(name=FONT, size=14, bold=True)
F_H = Font(name=FONT, size=16, bold=True)
FILL_H = PatternFill("solid", fgColor="D9E1F2")
FILL_T = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def interp_formula(ref):
    r = f"ROUND({ref},2)"
    return (f'=IF({ref}="","",IF({r}<=1.8,"น้อยที่สุด",IF({r}<=2.6,"น้อย",'
            f'IF({r}<=3.4,"ปานกลาง",IF({r}<=4.2,"มาก","มากที่สุด")))))')


def read_csv(p):
    if not Path(p).exists():
        return []
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.manifest = []

    def sheet(self, name):
        return self.wb.create_sheet(name[:31])

    def header(self, ws, row, labels, widths=None):
        for j, lab in enumerate(labels, 1):
            c = ws.cell(row, j, lab)
            c.font, c.fill, c.border, c.alignment = F_B, FILL_H, BOX, Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            for j, w in enumerate(widths, 1):
                ws.column_dimensions[COL(j)].width = w

    def put(self, ws, r, c, v, font=F_N, fill=None, fmt=None):
        cell = ws.cell(r, c, v)
        cell.font, cell.border, cell.alignment = font, BOX, WRAP
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        return cell


def raw_sheet(b, name, respondents, items, scores, text_of):
    """แผ่นข้อมูลดิบ: แถว = ผู้ตอบ, คอลัมน์ = รายการ คืนช่วงของแต่ละรายการ"""
    ws = b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    b.header(ws, 3, ["ผู้ตอบ"] + items, [12] + [9] * len(items))
    for j, it in enumerate(items, 2):
        ws.cell(2, j, text_of.get(it, "")[:60]).font = Font(name=FONT, size=10, italic=True)
    ranges = {}
    for i, resp in enumerate(respondents, 4):
        b.put(ws, i, 1, resp, F_IN)
        for j, it in enumerate(items, 2):
            v = scores.get(it, {}).get(resp, "")
            try:
                v = float(v) if v != "" and re.fullmatch(r"-?\d+(\.\d+)?", str(v)) else v
                v = int(v) if isinstance(v, float) and v.is_integer() else v
            except ValueError:
                pass
            b.put(ws, i, j, v, F_IN)
    last = 3 + len(respondents)
    for j, it in enumerate(items, 2):
        ranges[it] = f"'{ws.title}'!${COL(j)}$4:${COL(j)}${last}"
    ws.freeze_panes = "B4"
    return ranges


def score_sheet(b, name, groups, ranges, text_of, n_label="n"):
    """groups: OrderedDict(aspect -> [codes]) → รายข้อ + รายด้าน (pooled) + ภาพรวม"""
    ws = b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    b.header(ws, 3, ["รหัส", "รายการ", n_label, "X̄", "S.D.", "แปลผล"], [10, 60, 6, 9, 9, 14])
    r = 4
    all_codes = []
    for aspect, codes in groups.items():
        for c in codes:
            rg = ranges[c]
            b.put(ws, r, 1, c)
            b.put(ws, r, 2, text_of.get(c, ""))
            b.put(ws, r, 3, f"=COUNT({rg})")
            b.put(ws, r, 4, f"=AVERAGE({rg})", fmt="0.00")
            b.put(ws, r, 5, f"=_xlfn.STDEV.S({rg})", fmt="0.00")
            b.put(ws, r, 6, interp_formula(f"D{r}"))
            b.manifest.append({"sheet": ws.title, "row": r, "kind": "item", "items": [c]})
            r += 1
        all_codes += codes
        if aspect:
            rgs = ",".join(ranges[c] for c in codes)
            b.put(ws, r, 1, "", F_B, FILL_T)
            b.put(ws, r, 2, f"รวม{aspect}", F_B, FILL_T)
            b.put(ws, r, 3, f"=COUNT({rgs})", F_B, FILL_T)
            b.put(ws, r, 4, f"=AVERAGE({rgs})", F_B, FILL_T, "0.00")
            b.put(ws, r, 5, f"=_xlfn.STDEV.S({rgs})", F_B, FILL_T, "0.00")
            b.put(ws, r, 6, interp_formula(f"D{r}"), F_B, FILL_T)
            b.manifest.append({"sheet": ws.title, "row": r, "kind": "pooled", "items": codes, "label": aspect})
            r += 1
    rgs = ",".join(ranges[c] for c in all_codes)
    for j, v in enumerate(["", "ภาพรวม", f"=COUNT({rgs})", f"=AVERAGE({rgs})", f"=_xlfn.STDEV.S({rgs})", interp_formula(f"D{r}")], 1):
        b.put(ws, r, j, v, F_B, FILL_T, "0.00" if j in (4, 5) else None)
    b.manifest.append({"sheet": ws.title, "row": r, "kind": "pooled", "items": all_codes, "label": "ภาพรวม"})
    ws.cell(r + 2, 1, "หมายเหตุ: ค่ารายด้านและภาพรวมคำนวณจากคะแนนดิบทุกค่า S.D. แบบตัวอย่าง (n − 1) "
            "แปลผล 1.00–1.80 น้อยที่สุด 1.81–2.60 น้อย 2.61–3.40 ปานกลาง 3.41–4.20 มาก 4.21–5.00 มากที่สุด").font = F_N
    return ws, r


def choices_raw(b, choices):
    ws = b.sheet("ข้อมูลดิบ-ช่องเลือก")
    b.header(ws, 1, ["ผู้ตอบ", "คำถาม", "ตัวเลือก", "เลือก (1/0)"], [12, 50, 50, 10])
    for i, c in enumerate(choices, 2):
        for j, k in enumerate(["respondent", "question", "option", "checked"], 1):
            v = c[k]
            b.put(ws, i, j, int(v) if k == "checked" else v, F_IN)
    last = max(2, len(choices) + 1)
    return {"q": f"'{ws.title}'!$B$2:$B${last}", "o": f"'{ws.title}'!$C$2:$C${last}", "c": f"'{ws.title}'!$D$2:$D${last}"}


def choice_sheet(b, name, choices, rng, question_re, n_resp_ref, note=None):
    """จำนวนและร้อยละจากฐานผู้ตอบทั้งหมด"""
    ws = b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    if note:
        ws.cell(2, 1, note).font = F_N
    b.header(ws, 3, ["คำถาม", "ตัวเลือก", "จำนวน", "ร้อยละ"], [40, 55, 10, 10])
    seen = OrderedDict()
    for c in choices:
        if re.search(question_re, c["question"]):
            seen.setdefault((c["question"], c["option"]), None)
    r = 4
    for q, o in seen:
        b.put(ws, r, 1, q)
        b.put(ws, r, 2, o)
        b.put(ws, r, 3, f'=COUNTIFS({rng["q"]},A{r},{rng["o"]},B{r},{rng["c"]},1)')
        b.put(ws, r, 4, f"=IF({n_resp_ref}=0,0,C{r}/{n_resp_ref}*100)", fmt="0.00")
        b.manifest.append({"sheet": ws.title, "row": r, "kind": "choice", "question": q, "option": o})
        r += 1
    if r == 4:
        b.put(ws, r, 1, "(ไม่พบคำตอบที่ตรงกับคำถามนี้ — ตรวจ question pattern)")
    ws.cell(r + 1, 1, "ร้อยละคิดจากจำนวนผู้ตอบทั้งหมด (เลือกได้หลายข้อ ผลรวมอาจเกิน 100)").font = F_N
    return ws


def text_sheet(b, name, rows, cols=("respondent", "item", "column", "text"), labels=("ผู้ตอบ", "รหัส", "ช่อง", "ข้อความ"), extra=()):
    ws = b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    b.header(ws, 3, list(labels) + list(extra), [12, 10, 18, 80] + [20] * len(extra))
    for i, t in enumerate(rows, 4):
        for j, k in enumerate(cols, 1):
            b.put(ws, i, j, t.get(k, ""), F_IN)
        for j in range(len(extra)):
            b.put(ws, i, len(cols) + 1 + j, t.get(extra[j], ""), F_IN)
    if not rows:
        b.put(ws, 4, 1, "(ไม่มีข้อมูล)")
    return ws


def agree_sheet(b, name, items, ranges, k, levels, text_of, pct=0.8):
    """items: [code]; ranges[code#agree], ranges[code#moscow] (ถ้ามี); levels: code→ระดับเดิม"""
    ws = b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    need = L.pass_threshold(k, pct)
    ws.cell(2, 1, f"เกณฑ์ผ่าน: เห็นชอบตั้งแต่ร้อยละ {int(pct*100)} ({need} ใน {k} คน) — ไม่ผ่านให้ปรับเป็นระดับที่ถูกเสนอมากที่สุด").font = F_N
    has_lvl = any(levels.get(c) for c in items)
    hdr = ["รหัส", "รายการ", "เห็นชอบ", "ไม่เห็นชอบ", "ร้อยละเห็นชอบ", "ผล"]
    if has_lvl:
        hdr = ["รหัส", "รายการ", "ระดับเดิม", "เห็นชอบ", "ไม่เห็นชอบ", "ร้อยละเห็นชอบ", "เสนอ M", "เสนอ S", "เสนอ C", "เสนอ W",
               "ระดับที่ถูกเสนอมากที่สุด", "ผล", "ระดับฉบับสุดท้าย"]
    b.header(ws, 3, hdr, [10, 50] + [10] * (len(hdr) - 2))
    ws.cell(1, 20, k).font = F_IN
    ws.cell(1, 19, "จำนวนผู้เชี่ยวชาญ").font = F_N
    ws.cell(2, 20, need).font = F_N
    ws.cell(2, 19, "ขั้นต่ำที่ต้องเห็นชอบ").font = F_N
    K, NEED = "$T$1", "$T$2"
    for r, c in enumerate(items, 4):
        ag = ranges[f"{c}#agree"]
        b.put(ws, r, 1, c)
        b.put(ws, r, 2, text_of.get(c, ""))
        if not has_lvl:
            b.put(ws, r, 3, f'=COUNTIF({ag},"agree")')
            b.put(ws, r, 4, f'=COUNTIF({ag},"disagree")')
            b.put(ws, r, 5, f"=C{r}/{K}*100", fmt="0.00")
            b.put(ws, r, 6, f'=IF(C{r}>={NEED},"ถึงเกณฑ์ ("&C{r}&"/"&{K}&")","ไม่ถึงเกณฑ์ ("&C{r}&"/"&{K}&")")')
            b.manifest.append({"sheet": ws.title, "row": r, "kind": "agree", "items": [c], "agree_col": 3})
            continue
        mo = ranges.get(f"{c}#moscow")
        b.put(ws, r, 3, levels.get(c, ""), F_IN)
        b.put(ws, r, 4, f'=COUNTIF({ag},"agree")')
        b.put(ws, r, 5, f'=COUNTIF({ag},"disagree")')
        b.put(ws, r, 6, f"=D{r}/{K}*100", fmt="0.00")
        for j, lv in enumerate("MSCW", 7):
            b.put(ws, r, j, f'=COUNTIF({mo},"{lv}")' if mo else 0)
        b.put(ws, r, 11, f'=IF(SUM(G{r}:J{r})=0,"",IF(COUNTIF(G{r}:J{r},MAX(G{r}:J{r}))>1,"เสมอ — ผู้วิจัยตัดสินพร้อมเหตุผล",'
                         f'INDEX({{"M","S","C","W"}},MATCH(MAX(G{r}:J{r}),G{r}:J{r},0))))')
        b.put(ws, r, 12, f'=IF(D{r}>={NEED},"ถึงเกณฑ์ ("&D{r}&"/"&{K}&")","ปรับตามเสียงข้างมาก ("&D{r}&"/"&{K}&")")')
        b.put(ws, r, 13, f'=IF(D{r}>={NEED},C{r},K{r})', F_B)
        b.manifest.append({"sheet": ws.title, "row": r, "kind": "agree", "items": [c], "agree_col": 4})
    return ws


def info_sheet(b, name, choices, texts, rng, info_re, n_ref):
    ws = choice_sheet(b, name, choices, rng, info_re, n_ref,
                      note="ข้อมูลทั่วไปของผู้ตอบ (คำถามที่เลือกได้หลายข้อ ร้อยละคิดจากผู้ตอบทั้งหมด)")
    return ws


def summary_sheet(b, name, lines, first=False):
    ws = b.wb.create_sheet(name, 0) if first else b.sheet(name)
    ws.cell(1, 1, name).font = F_H
    b.header(ws, 3, ["ประเด็น", "ค่า", "หมายเหตุ"], [45, 18, 60])
    for i, (lab, val, note) in enumerate(lines, 4):
        b.put(ws, i, 1, lab)
        b.put(ws, i, 2, val, fmt="0.00" if isinstance(val, str) and "AVERAGE" in val else None)
        b.put(ws, i, 3, note)
    return ws


def respondents_of(scores, choices, texts):
    s = OrderedDict()
    for src in (choices, texts):
        for r in src:
            s.setdefault(r["respondent"], None)
    for it in scores.values():
        for r in it:
            s.setdefault(r, None)
    return sorted(s)


def build(mode, proj, data_dir, out):
    scores = L.load_scores(Path(data_dir) / "scores.csv")
    choices = read_csv(Path(data_dir) / "choices.csv")
    texts = read_csv(Path(data_dir) / "text.csv")
    resp = respondents_of(scores, choices, texts)
    text_of = {x["code"]: x["text"] for k in ("problems", "needs", "requirements", "arch_items", "eval_items", "confirm_items")
               for x in proj.get(k, [])}
    b = Book()
    k = len(resp)
    if mode == "user":
        dims = OrderedDict()
        dims["ปัญหา (P)"] = [p["code"] for p in proj["problems"]]
        dim_name = {d["code"]: d["name"] for d in proj["dimensions"]}
        for n in proj["needs"]:
            dims.setdefault(dim_name[n["dimension"]], []).append(n["code"])
        items = [c for cs in dims.values() for c in cs]
        missing = [c for c in items if c not in scores]
        if missing:
            print("⚠ ไม่มีคะแนนของ:", missing)
        rng = choices_raw(b, choices)
        ranges = raw_sheet(b, "ข้อมูลดิบ", resp, items, scores, text_of)
        nref = f"'สรุปผล'!$B$4"
        info_sheet(b, "ข้อมูลทั่วไป", choices, texts, rng, proj.get("info_question_regex", "บทบาท|ศูนย์|หน่วยงาน|ประสบการณ์|ความถี่"), nref)
        ws, last = score_sheet(b, "ค่าเฉลี่ยรายด้าน", dims, ranges, text_of)
        choice_sheet(b, "การเลือกฟังก์ชัน", choices, rng, proj.get("function_question_regex", "ฟังก์ชัน|พัฒนาก่อน"), nref)
        text_sheet(b, "คำตอบปลายเปิด", [t for t in texts if t["item"].startswith("OP")])
        need_rows = [m for m in b.manifest if m["sheet"] == "ค่าเฉลี่ยรายด้าน" and m["kind"] == "pooled"]
        lines = [("จำนวนผู้ตอบ (n)", k, "นับจากไฟล์แบบสอบถามที่อ่านได้")]
        for m in need_rows:
            lines.append((f"X̄ {m['label']}", f"='ค่าเฉลี่ยรายด้าน'!D{m['row']}", f"S.D. ='ค่าเฉลี่ยรายด้าน'!E{m['row']}"))
        summary_sheet(b, "สรุปผล", lines, first=True)
    elif mode == "expert1":
        musts = proj["round1"]["score_items"]
        confirm = proj["round1"]["confirm_items"]
        levels = {c: lv for c, lv in proj["round1"].get("levels", {}).items()}
        rng = choices_raw(b, choices)
        nref = "'สรุปผล'!$B$4"
        info_sheet(b, "ข้อมูลผู้เชี่ยวชาญ", choices, texts, rng, proj.get("expert_info_regex", "ตำแหน่ง|ศูนย์|วุฒิ|ประสบการณ์|งานที่"), nref)
        items = musts + [f"{c}#agree" for c in confirm] + [f"{c}#moscow" for c in confirm]
        ranges = raw_sheet(b, "ข้อมูลดิบ", resp, items, scores, text_of)
        score_sheet(b, "ค่าเฉลี่ย", OrderedDict([("ข้อ Must", musts)]), ranges, text_of)
        agree_sheet(b, "ยืนยัน MoSCoW", confirm, ranges, k, levels, text_of)
        choice_sheet(b, "ความเห็นโดยรวม", choices, rng, proj.get("overall_question_regex", "ความเห็นโดยรวม"), nref)
        text_sheet(b, "เหตุผลรายข้อ", [t for t in texts if not t["item"].startswith("OP")])
        text_sheet(b, "คำตอบปลายเปิด", [t for t in texts if t["item"].startswith("OP")])
        synth = proj["round1"].get("synthesis", [])
        text_sheet(b, "สังเคราะห์ประเด็น", synth, cols=("issue", "experts", "decision", "applied_to"),
                   labels=("ประเด็น", "ผู้เสนอ", "การตัดสินใจ", "นำไปปรับที่"))
        summary_sheet(b, "สรุปผล", [("จำนวนผู้เชี่ยวชาญ", k, "ชุดเดียวกันตลอดงาน"),
                                     ("X̄ ภาพรวมข้อ Must", f"='ค่าเฉลี่ย'!D{4 + len(musts) + 1}", "จากคะแนนดิบทุกค่า"),
                                     ("เกณฑ์ผ่าน", f"{L.pass_threshold(k)}/{k}", "ร้อยละ 80")], first=True)
    elif mode == "expert2":
        r2 = proj["round2"]
        arch = OrderedDict()
        for it in proj["arch_items"]:
            arch.setdefault(it["aspect"], []).append(it["code"])
        ev = OrderedDict()
        for it in proj["eval_items"]:
            ev.setdefault(it["aspect"], []).append(it["code"])
        confirm = [c["code"] for c in proj["confirm_items"]]
        rng = choices_raw(b, choices)
        nref = "'สรุปผล'!$B$4"
        info_sheet(b, "ข้อมูลผู้เชี่ยวชาญ", choices, texts, rng, proj.get("expert_info_regex", "ตำแหน่ง|ศูนย์|วุฒิ|ประสบการณ์|งานที่"), nref)
        items = [c for cs in arch.values() for c in cs] + [c for cs in ev.values() for c in cs] + [f"{c}#agree" for c in confirm]
        ranges = raw_sheet(b, "ข้อมูลดิบ-คะแนน", resp, items, scores, text_of)
        wa, la = score_sheet(b, "ค่าเฉลี่ย-สถาปัตยกรรม", arch, ranges, text_of)
        we, le = score_sheet(b, "ค่าเฉลี่ย-ข้อกำหนด", ev, ranges, text_of)
        agree_sheet(b, "การยืนยัน", confirm, ranges, k, {}, text_of)
        choice_sheet(b, "ความเห็นโดยรวม", choices, rng, proj.get("overall_question_regex", "ความเห็นโดยรวม"), nref)
        sugg = r2.get("suggestions") or [{"respondent": t["respondent"], "item": t["item"], "column": t.get("column", ""), "text": t["text"], "category": ""} for t in texts]
        ws = text_sheet(b, "ข้อเสนอแนะ", sugg, extra=("category", "action"))
        ws.cell(2, 1, "category: ชื่นชม/ยืนยัน หรือ เสนอเพิ่ม/ปรับ — action: สิ่งที่ปรับ (ระบุ 'เพิ่มหลังการประเมิน')").font = F_N
        text_sheet(b, "สังเคราะห์ประเด็น", r2.get("synthesis", []), cols=("issue", "experts", "decision", "applied_to"),
                   labels=("ประเด็น", "ผู้เสนอ", "การตัดสินใจ", "นำไปปรับที่"))
        summary_sheet(b, "สรุปผล", [("จำนวนผู้เชี่ยวชาญ", k, "ชุดเดียวกับครั้งที่ 1"),
                                     ("X̄ ภาพรวมสถาปัตยกรรม", f"='ค่าเฉลี่ย-สถาปัตยกรรม'!D{la}", "จากคะแนนดิบทุกค่า"),
                                     ("S.D. ภาพรวมสถาปัตยกรรม", f"='ค่าเฉลี่ย-สถาปัตยกรรม'!E{la}", ""),
                                     ("X̄ ภาพรวมข้อกำหนด", f"='ค่าเฉลี่ย-ข้อกำหนด'!D{le}", "จากคะแนนดิบทุกค่า"),
                                     ("S.D. ภาพรวมข้อกำหนด", f"='ค่าเฉลี่ย-ข้อกำหนด'!E{le}", ""),
                                     ("เกณฑ์ผ่านการยืนยัน", f"{L.pass_threshold(k)}/{k}", "ร้อยละ 80")], first=True)
    else:
        raise SystemExit("mode ต้องเป็น user/expert1/expert2")
    for ws in b.wb.worksheets:
        ws.sheet_view.zoomScale = 100
    b.wb.save(out)
    Path(str(out) + ".manifest.json").write_text(json.dumps({"mode": mode, "data": str(data_dir), "k": k, "rows": b.manifest},
                                                            ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"บันทึก {out}: ผู้ตอบ {k} ราย, แผ่น {[ws.title for ws in b.wb.worksheets]}")
    print("→ ตรวจต่อ: python verify_excel.py", out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mode", required=True, choices=("user", "expert1", "expert2"))
    ap.add_argument("--project", required=True)
    ap.add_argument("--data", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    build(a.mode, L.load_project(a.project), a.data, a.out)


if __name__ == "__main__":
    main()
